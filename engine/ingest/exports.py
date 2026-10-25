"""Transcript exports: paragraphs by speaker with a timestamp on every paragraph."""
from __future__ import annotations

from pathlib import Path
from typing import Any


def _ts(s: float) -> str:
    s = max(0, int(s))
    return f"{s // 3600:02d}:{(s % 3600) // 60:02d}:{s % 60:02d}"


def paragraphs(words: list[dict[str, Any]], speakers: dict[str, str], max_len_s: float = 60.0,
               skip_deleted: bool = False) -> list[dict[str, Any]]:
    """Group words into paragraphs: new paragraph on speaker change, or every ~60 s at a
    sentence end.  Each paragraph: {speaker, name, start, end, text}."""
    paras: list[dict[str, Any]] = []
    cur: list[dict[str, Any]] = []

    def flush():
        if cur:
            paras.append({"speaker": cur[0]["speaker"], "name": speakers.get(cur[0]["speaker"], cur[0]["speaker"]),
                          "start": cur[0]["start"], "end": cur[-1]["end"],
                          "text": " ".join(w["text"] if "text" in w else w["word"] for w in cur)})
            cur.clear()

    for w in words:
        if skip_deleted and w.get("deleted"):
            continue
        if cur and (w["speaker"] != cur[0]["speaker"]
                    or (w["start"] - cur[0]["start"] > max_len_s and (cur[-1].get("text") or cur[-1].get("word", "")).endswith((".", "?", "!")))):
            flush()
        cur.append(w)
    flush()
    return paras


def write_txt(paras: list[dict[str, Any]], path: Path, title: str = "") -> None:
    lines = [title, ""] if title else []
    for p in paras:
        lines.append(f"[{_ts(p['start'])}] {p['name']}:")
        lines.append(p["text"])
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def write_docx(paras: list[dict[str, Any]], path: Path, title: str, meta: dict[str, Any] | None = None) -> None:
    from docx import Document
    from docx.shared import Pt, RGBColor
    doc = Document()
    doc.add_heading(title or "Transcript", level=1)
    if meta:
        p = doc.add_paragraph()
        for k, v in meta.items():
            if v:
                r = p.add_run(f"{k}: {v}\n"); r.font.size = Pt(10); r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    for para in paras:
        p = doc.add_paragraph()
        r = p.add_run(f"[{_ts(para['start'])}] "); r.font.color.rgb = RGBColor(0x88, 0x88, 0x88); r.font.size = Pt(9)
        r = p.add_run(f"{para['name']}: "); r.bold = True
        p.add_run(para["text"])
    doc.save(str(path))
